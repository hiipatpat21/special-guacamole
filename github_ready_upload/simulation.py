"""Student allowance simulation: Monday to Thursday, four budgeting scenarios.

Run this file with Python. Edit TRIALS and SEED below to change the experiment.
All monetary calculations use integer centavos; outputs use Philippine pesos.
"""

from decimal import Decimal, InvalidOperation
from pathlib import Path
import random
import warnings

from openpyxl import load_workbook
from analysis import add_intervals, comparison_row, empty_pair, export_facts, record_pair, save_csv
from visuals import save_visuals


FOLDER = Path(__file__).resolve().parent
WORKBOOK = FOLDER / "data" / "Data_20-students.xlsx"
TRIALS = 10_000  # per student, per scenario: 20 x 10,000 x 4 = 800,000 trials
SEED = 20261027
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday"]
INPUT_COLUMNS = {"allowance": 3, "food": 4, "transport": 5, "optional": 6,
                 "supplies": 7, "event_count": 10, "event_cost": 11}
# Scenario: allowance %, optional spending %, protected reserve %.
SCENARIOS = {
    "Normal spending": (100, 100, 0),
    "Lower allowance": (80, 100, 0),
    "Controlled spending": (100, 50, 0),
    "Emergency reserve": (100, 100, 20),
}


def cents(value):
    """Read a nonnegative peso amount with at most two decimal places (rule 15)."""
    try:
        amount = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError("Money must be a number.") from error
    if not amount.is_finite() or amount < 0 or amount * 100 != (amount * 100).to_integral():
        raise ValueError("Money must be nonnegative with at most two decimals.")
    return int(amount * 100)


def percent(amount, rate):
    """Multiply integer centavos by a percentage, rounding half a centavo up."""
    return (amount * rate + 50) // 100


def validate_student(student):
    """Rule 15: validate the model inputs once, before the repeated trials."""
    for field in INPUT_COLUMNS:
        if not isinstance(student[field], int) or isinstance(student[field], bool) or student[field] < 0:
            raise ValueError(f"{field} must be a nonnegative integer (money in centavos).")
    if student["event_count"] > 16:
        raise ValueError("Event count above 16 requires a different event model.")
    if student["event_count"] > 0 and student["event_cost"] == 0:
        raise ValueError("A positive event count needs a positive event cost.")


def load_students():
    """Keep each workbook row together; reserves are already part of allowance."""
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    students = []
    try:
        rows = workbook["Modified Sample Data"].iter_rows(values_only=True)
        headers = next(rows)
        for column, word in {3: "money", 4: "food", 5: "transportation", 6: "optional", 7: "printing", 10: "unexpected", 11: "unexpected"}.items():
            if word not in str(headers[column]).lower():
                raise ValueError("The workbook columns do not match the questionnaire.")
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None for value in row):
                continue
            if len(row) < 16 or any(value is None for value in row[1:16]):
                raise ValueError(f"Row {row_number}: a questionnaire field is missing.")
            if str(row[1]).strip().lower() != "yes":
                raise ValueError(f"Row {row_number}: student must attend all four campus days.")
            count = Decimal(str(row[10]))
            if not count.is_finite() or count != count.to_integral() or not 0 <= count <= 16:
                raise ValueError(f"Row {row_number}: event count must be an integer from 0 to 16.")
            student = {"id": f"S{row_number - 1:02d}", "row": row_number, "event_count": int(count)}
            for field, column in INPUT_COLUMNS.items():
                if field != "event_count":
                    student[field] = cents(row[column])
            validate_student(student)
            low, high, reported_reserve = cents(row[8]), cents(row[9]), cents(row[14])
            if low > high or reported_reserve > student["allowance"] or (str(row[13]).lower() == "no" and reported_reserve > 0):
                warnings.warn(f"{student['id']}: review the reported low/high or reserve values; inputs retained.")
            students.append(student)
    finally:
        workbook.close()
    if not students:
        raise ValueError("The workbook contains no students.")
    return students


def spending_multiplier(uniform):
    """Rules 3, 5-6: use exact uniform thresholds, without rounding probabilities."""
    if uniform < 0.25:
        return 80
    if uniform < 0.75:
        return 100
    return 120


def pay_day(balance, reserve, required, desired_optional):
    """Rules 7-10, 13-14: essentials first; optional costs use ordinary cash only."""
    if required > balance:
        return 0, 0, 0, required - balance  # first shortfall; consume available cash
    balance -= required
    reserve = min(reserve, balance)  # required costs may use reserve; never refill it
    optional_paid = min(desired_optional, balance - reserve)
    return balance - optional_paid, reserve, optional_paid, 0


def simulate_week(student, scenario, draws, trace=None):
    """One student's four days; draws contains all four days before any failure."""
    allowance_rate, optional_rate, reserve_rate = SCENARIOS[scenario]
    balance = percent(student["allowance"], allowance_rate)  # rules 1, 11
    reserve = percent(balance, reserve_rate)  # part of balance, never extra money
    supplies, remainder = divmod(student["supplies"], 4)  # rule 4: exact total
    forgone = desired_total = 0
    shortage_day = shortfall = 0
    for day, (spending_draw, event_draw) in enumerate(draws):
        multiplier = spending_multiplier(spending_draw)
        emergency = student["event_cost"] if event_draw < student["event_count"] / 16 else 0
        required = percent(student["food"], multiplier) + student["transport"]
        required += supplies + (day < remainder) + emergency
        # Rule 12: combine multiplier and optional reduction before rounding once.
        desired = (student["optional"] * multiplier * optional_rate + 5000) // 10000
        cash_before = balance
        balance, reserve, paid, shortfall = pay_day(balance, reserve, required, desired)
        if trace is not None:
            trace.append([DAYS[day], cash_before / 100, required / 100, desired / 100,
                          paid / 100, reserve / 100, balance / 100])
        desired_total += desired
        forgone += desired - paid
        if shortfall:
            shortage_day = day + 1
            break  # no demand is counted after the first shortage day
    return {"day": shortage_day, "cash": balance, "shortfall": shortfall,
            "forgone": forgone, "desired": desired_total}


def empty_totals():
    return {"days": [0] * 5, "cash": 0, "shortfall": 0, "forgone": 0, "desired": 0,
            "cash_squared": 0}


def record(totals, outcome):
    totals["days"][outcome["day"]] += 1  # index 0 means no shortage
    for field in ("cash", "shortfall", "forgone", "desired"):
        totals[field] += outcome[field]
    totals["cash_squared"] += outcome["cash"] ** 2


def result_row(student_id, scenario, totals, profile_totals=None):
    days = totals["days"]
    trials, failures = sum(days), sum(days[1:])
    mode = " / ".join(DAYS[i] for i in range(4) if days[i + 1] == max(days[1:])) if failures else "no shortages"
    row = {"student": student_id, "scenario": scenario, "trials": trials,
           "shortage_probability": failures / trials,
           "mean_thursday_cash_php": totals["cash"] / trials / 100,
           "most_frequent_first_shortage_day": mode,
           "mean_cash_successful_trials_php": totals["cash"] / days[0] / 100 if days[0] else None,
           "mean_first_shortfall_failed_trials_php": totals["shortfall"] / failures / 100 if failures else None,
           "mean_optional_forgone_reached_days_php": totals["forgone"] / trials / 100,
           "mean_optional_desired_reached_days_php": totals["desired"] / trials / 100,
           "no_shortage_probability": days[0] / trials}
    for i, day in enumerate(DAYS, start=1):
        row[f"first_shortage_{day.lower()}_probability"] = days[i] / trials
    return add_intervals(row, profile_totals if profile_totals is not None else [totals])


def main():
    if not isinstance(TRIALS, int) or TRIALS < 2:
        raise ValueError("TRIALS must be an integer of at least 2.")
    students = load_students()
    overall = {name: empty_totals() for name in SCENARIOS}
    profile_totals = {name: [] for name in SCENARIOS}
    paired_profiles = {name: [] for name in list(SCENARIOS)[1:]}
    rows, comparisons = [], []
    for student in students:
        rng = random.Random(SEED + student["row"])
        totals = {name: empty_totals() for name in SCENARIOS}
        paired = {name: empty_pair() for name in paired_profiles}
        for _ in range(TRIALS):
            # Reuse identical four-day draws for every scenario, even after failure.
            draws = [(rng.random(), rng.random()) for _ in DAYS]
            outcomes = {name: simulate_week(student, name, draws) for name in SCENARIOS}
            for scenario, outcome in outcomes.items():
                record(totals[scenario], outcome)
                record(overall[scenario], outcome)
                if scenario in paired:
                    record_pair(paired[scenario], outcomes["Normal spending"], outcome)
        rows.extend(result_row(student["id"], name, totals[name]) for name in SCENARIOS)
        for name in SCENARIOS:
            profile_totals[name].append(totals[name])
        for name in paired:
            paired_profiles[name].append(paired[name])
            comparisons.append(comparison_row(student["id"], name, [paired[name]]))
    # Every profile has the same trial count, so pooled averages give equal weights.
    summary = [result_row("ALL_EQUAL_WEIGHT", name, overall[name], profile_totals[name]) for name in SCENARIOS]
    paired_summary = [comparison_row("ALL_EQUAL_WEIGHT", name, profiles) for name, profiles in paired_profiles.items()]
    output = FOLDER / "results"
    output.mkdir(exist_ok=True)
    save_csv(output / "results.csv", summary + rows)
    save_csv(output / "comparisons.csv", paired_summary + comparisons)
    fact_count = export_facts(students, INPUT_COLUMNS, output)
    # Verification case 9: fixed teaching example, separate from the random trials.
    example = {"allowance": 10000, "food": 1500, "transport": 0, "optional": 2000,
               "supplies": 0, "event_count": 0, "event_cost": 0}
    trace = []
    simulate_week(example, "Emergency reserve", [(0.5, 0.99)] * 4, trace)
    save_visuals(summary, output, len(students), TRIALS, trace)
    (output / "run_settings.txt").write_text(
        f"Seed: {SEED}\nGenerator: Python random.Random(seed + workbook row)\n"
        f"Profiles: {len(students)}\nTrials per profile per scenario: {TRIALS}\n"
        f"Total main trials: {len(students) * TRIALS * len(SCENARIOS)}\n"
        f"Input facts: {fact_count}\nInputs: AI-adjusted survey values\n"
        "Intervals: normal 95% Monte Carlo intervals, fixed-profile weights, paired differences.\n",
        encoding="utf-8")
    print(f"Seed: {SEED}; {len(students)} profiles; {TRIALS:,} trials per profile per scenario.")
    print("AI-adjusted survey inputs; results describe these profiles under the model assumptions.\n")
    print(f"{'Scenario':22} {'Shortage':>10} {'Mean cash (PHP)':>18}")
    for row in summary:
        print(f"{row['scenario']:22} {row['shortage_probability']:10.2%} {row['mean_thursday_cash_php']:18.2f}")
    print(f"\nSaved results with 95% Monte Carlo intervals, paired comparisons, {fact_count} input facts,")
    print(f"three charts, a worked example, and run settings to {output}")


if __name__ == "__main__":
    main()

